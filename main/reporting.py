"""Summarize simulation operating characteristics and paired Monte Carlo differences."""
from __future__ import annotations
import csv, json, math
from collections import Counter
from pathlib import Path
import numpy as np
from analysis import METHOD_DESCRIPTIONS
from reference.core import signature, canonical_labels, adjusted_rand_value, exact_truth, OBS_PROB
from cluster_selection import ORDINAL

Z95 = 1.959963984540054


def mean_se(x):
    x = np.asarray(x, float)
    return float(x.mean()), float(x.std(ddof=1)/np.sqrt(len(x)))


def rate(x):
    x = np.asarray(x, bool); R = len(x); h = int(x.sum()); p = h/R
    zz = Z95**2; den = 1+zz/R
    mid = (p+zz/(2*R))/den
    half = Z95*np.sqrt(p*(1-p)/R+zz/(4*R**2))/den
    return dict(hits=h, rate=p, mcse=float(np.sqrt(p*(1-p)/R)), lo95=max(0.,mid-half), hi95=min(1.,mid+half))


def metric_fields(name, x):
    rr = rate(x)
    return {name+'_'+key: value for key,value in rr.items()}


def status(p, p0):
    if abs(p-p0) <= 1e-10: return 'null_boundary'
    return 'null_interior' if p < p0 else 'alternative'


def keys(e):
    return dict(setting=e.name, family=e.family, scenario=e.scenario, n=e.n, dynamics=e.dynamics)


def write_csv(path, rows):
    if not rows: return
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(dict.fromkeys(k for r in rows for k in r))
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.DictWriter(f, fieldnames=fields); writer.writeheader(); writer.writerows(rows)


def partition_metrics(z, truth):
    """Labels have already been returned to original basket order."""
    ii,jj = np.triu_indices(5,1)
    predicted = z[:,ii] == z[:,jj]; correct = truth[ii] == truth[jj]
    k = np.array([len(np.unique(x)) for x in z])
    K = len(np.unique(truth))
    exact = np.all(predicted == correct, axis=1)
    ari = np.array([adjusted_rand_value(truth,x) for x in z])
    false_merge = np.mean(predicted[:,~correct],axis=1) if np.any(~correct) else np.full(len(z),np.nan)
    false_split = np.mean(~predicted[:,correct],axis=1) if np.any(correct) else np.full(len(z),np.nan)
    return dict(k=k, correct_number=k==K, correct_partition=exact, ARI=ari,
                under=k<K, over=k>K, false_merge=false_merge, false_split=false_split)


def paired_fields(x,y):
    d=np.asarray(x,float)-np.asarray(y,float)
    if not np.isfinite(d).all():
        return dict(difference=None,mcse=None,lo95=None,hi95=None)
    avg,se=mean_se(d)
    return dict(difference=avg,mcse=se,lo95=avg-Z95*se,hi95=avg+Z95*se)


def load_experiment(out,e,config,fingerprint):
    from simulate import chunk_path, check_chunk
    blocks=[]
    # Patient-level arrays remain in checkpoints; summary needs basket totals only.
    names=('post','labels','responses','basket_order','initial_counts','transition_counts',
           'final_counts','length_counts','Phat','final_feature','occupancy_feature','fitted_hitting','replicate')
    for start in range(0,config['reps'],config['chunk_size']):
        end=min(start+config['chunk_size'],config['reps']); path=chunk_path(out,e,start,end)
        if not path.exists(): raise FileNotFoundError(f'Missing chunk: {path}')
        check_chunk(path,fingerprint,start,end,config['methods'])
        with np.load(path,allow_pickle=False) as d: blocks.append({k:d[k] for k in names})
    a={k:np.concatenate([x[k] for x in blocks]) for k in names}
    if not np.array_equal(a['replicate'],np.arange(config['reps'])):
        raise ValueError('Replicates missing, duplicated or out of order.')
    return a


def summarize(out,exps,config,manifest,excel=True):
    out=Path(out); methods=config['methods']; R=config['reps']
    names=('truth','cluster_metrics','cluster_patterns','basket_metrics','trial_metrics',
           'paired_cluster','paired_efficacy','dynamics_comparison','refinement_audit',
           'feature_diagnostics','transition_matrices','schedule','methods')
    tables={n:[] for n in names}; baseline_time={}
    for m in methods:
        tables['methods'].append(dict(method=m,description=METHOD_DESCRIPTIONS[m],
                       primary=(m=='Proposed'),source='Frozen modules and METHODS.md in this package'))
    tables['schedule']=[dict(assessments=t,probability=float(w),source='Original supplied simulator OBS_PROB')
                        for t,w in enumerate(OBS_PROB,1)]
    for e in exps:
        a=load_experiment(out,e,config,manifest['fingerprint']); tk=keys(e); post=a['post']
        true_p=e.true_orr; null=true_p<=e.p0+1e-10; truth=e.labels
        K=len(np.unique(truth)); part_stats={}
        for j in range(5):
            t=exact_truth(e.initial[j],e.sequences[j]); marg=t['marginals']
            true_occ=np.sum(OBS_PROB[:,None]*np.cumsum(marg,axis=0)/np.arange(1,11)[:,None],axis=0)
            tables['truth'].append({**tk,'basket':j+1,'true_group':int(truth[j])+1,'true_K':K,
               'true_ORR':float(true_p[j]),'null_threshold':e.p0,'legacy_threshold':e.legacy_p0,
               'efficacy_status':status(true_p[j],e.p0),'source':e.source,
               'early_SD_response_odds_multiplier':float(e.early_factors[j]) if e.early_factors is not None else None,
               'late_nonPD_to_PD_odds_multiplier':e.late_odds,
               **{f'initial_{s}':float(e.initial[j,l]) for l,s in enumerate(('CR','PR','SD','PD'))}})
            for l in range(9):
                for r,s0 in enumerate(('CR','PR','SD','PD')):
                    tables['transition_matrices'].append({**tk,'basket':j+1,'from_assessment':l+1,
                       'to_assessment':l+2,'from_state':s0,
                       **{f'to_{s}':float(e.sequences[j,l,r,k]) for k,s in enumerate(('CR','PR','SD','PD'))}})
            for k,s in enumerate(('CR','PR','SD','PD')):
                oc=a['occupancy_feature'][:,j,k]; ff=a['final_feature'][:,j,k]
                tables['feature_diagnostics'].append({**tk,'basket':j+1,'state':s,
                  'true_mean_occupancy':float(true_occ[k]),'mean_fitted_occupancy':float(oc.mean()),
                  'occupancy_MSE_vs_true':float(np.mean((oc-true_occ[k])**2)),
                  'true_final_probability':float(t['final'][k]),'mean_fitted_final':float(ff.mean()),
                  'final_MSE_vs_true':float(np.mean((ff-t['final'][k])**2)),
                  'zero_outgoing_rate':float(np.mean(a['transition_counts'][:,j,k].sum(-1)==0)),
                  'mean_outgoing':float(a['transition_counts'][:,j,k].sum(-1).mean()),
                  'transition_MSE_vs_pooled_limit':float(np.mean((a['Phat'][:,j,k]-t['pooled_limit'][k])**2)),
                  'true_ORR':float(true_p[j]),'mean_fitted_hitting':float(a['fitted_hitting'][:,j].mean()),
                  'hitting_MSE_vs_true':float(np.mean((a['fitted_hitting'][:,j]-true_p[j])**2))})
        for m,method in enumerate(methods):
            if method!='EXNEX':
                lab=a['labels'][:,m]
                metric=partition_metrics(lab,truth); part_stats[method]=metric
                row={**tk,'method':method,'reps':R,'true_K':K,'true_partition':signature(truth)}
                for name in ('correct_number','correct_partition','under','over'):
                    row.update(metric_fields(name,metric[name]))
                for name in ('ARI','false_merge','false_split'):
                    arr=metric[name]
                    avg,se=mean_se(arr) if np.isfinite(arr).all() else (None,None)
                    row[name+'_mean']=avg;row[name+'_mcse']=se
                row.update({f'prob_K{k}':float(np.mean(metric['k']==k)) for k in range(1,6)})
                tables['cluster_metrics'].append(row)
                cnt=Counter(signature(x) for x in lab)
                for sig,h in sorted(cnt.items(),key=lambda kv:(-kv[1],kv[0])):
                    tables['cluster_patterns'].append({**tk,'method':method,'partition':sig,'count':h,'reps':R,'probability':h/R})
            if config['cluster_only']: continue
            v=post[:,m]; rejects=v[:,:,1]>e.p0
            for j,p in enumerate(true_p):
                err=v[:,j,0]-p; cover=(v[:,j,1]<=p)&(p<=v[:,j,2])
                bias,bse=mean_se(err); mse,mse_se=mean_se(err**2); width,wse=mean_se(v[:,j,2]-v[:,j,1])
                row={**tk,'method':method,'basket':j+1,'reps':R,'true_ORR':float(p),'p0':e.p0,
                     'efficacy_status':status(p,e.p0),
                     'decision_interpretation':'type_I_error' if null[j] else 'power',
                     **metric_fields('rejection',rejects[:,j]),
                     'rejection_at_legacy_threshold':float(np.mean(v[:,j,1]>e.legacy_p0)),
                     'observed_ORR_mean':float(a['responses'][:,j].mean()/e.n),
                     'posterior_mean':float(v[:,j,0].mean()),'bias':bias,'bias_mcse':bse,
                     'MSE':mse,'MSE_mcse':mse_se,'RMSE':math.sqrt(mse),
                     **metric_fields('coverage90',cover),'mean_lower90':float(v[:,j,1].mean()),
                     'mean_upper90':float(v[:,j,2].mean()),'mean_width90':width,'width_mcse':wse,
                     'mean_posterior_prob_gt_null':float(v[:,j,3].mean()),
                     'mean_posterior_sd':float(v[:,j,5].mean()),
                     'mean_posterior_EX_probability':float(v[:,j,6].mean()) if method=='EXNEX' else None,
                     'CDF_CI_disagreements':int(np.sum(rejects[:,j]!=(v[:,j,3]>.95)))}
                tables['basket_metrics'].append(row)
            tr={**tk,'method':method,'reps':R,'null_baskets':int(null.sum()),
                 'active_baskets':int((~null).sum()),'error_control':'Not calibrated to a nominal frequentist level'}
            if null.any():tr.update(metric_fields('FWER',np.any(rejects[:,null],axis=1)))
            else:tr.update({f'FWER_{n}':None for n in ('hits','rate','mcse','lo95','hi95')})
            if (~null).any():tr.update(metric_fields('all_active_detected',np.all(rejects[:,~null],axis=1)))
            tables['trial_metrics'].append(tr)
        if 'Proposed' in part_stats and 'Previous' in part_stats:
            ip=methods.index('Proposed'); ib=methods.index('Previous')
            zp=a['labels'][:,ip];zb=a['labels'][:,ib];ii,jj=np.triu_indices(5,1)
            changed=np.any((zp[:,ii]==zp[:,jj])!=(zb[:,ii]==zb[:,jj]),axis=1)
            b=part_stats['Previous'];p=part_stats['Proposed']
            k_changes=int(np.sum(b['k']!=p['k']))
            non2_changes=int(np.sum(changed&(b['k']!=2)))
            exact_changes=int(np.sum(b['correct_partition']!=p['correct_partition']))
            if k_changes or non2_changes or (K!=2 and exact_changes):raise AssertionError('Frozen refinement invariance failure.')
            tables['refinement_audit'].append({**tk,'reps':R,'K_changes':k_changes,'non2_membership_changes':non2_changes,
                'membership_changes':int(changed.sum()),'exact_changes':exact_changes,
                'repaired':int(np.sum(~b['correct_partition']&p['correct_partition'])),
                'spoiled':int(np.sum(b['correct_partition']&~p['correct_partition'])),
                'wrong_to_wrong':int(np.sum(changed&~b['correct_partition']&~p['correct_partition']))})
        if 'Proposed' in methods:
            ip=methods.index('Proposed')
            for m,method in enumerate(methods):
                if m==ip: continue
                if method in part_stats:
                    for name in ('correct_number','correct_partition','ARI','false_merge','false_split'):
                        tables['paired_cluster'].append({**tk,'comparison':'Proposed minus '+method,'metric':name,
                            **paired_fields(part_stats['Proposed'][name],part_stats[method][name]),'reps':R})
                if config['cluster_only']:continue
                for j,p in enumerate(true_p):
                    ap=post[:,ip,j];bp=post[:,m,j]
                    measures={'rejection':(ap[:,1]>e.p0,bp[:,1]>e.p0),
                              'squared_error':((ap[:,0]-p)**2,(bp[:,0]-p)**2),
                              'coverage90':((ap[:,1]<=p)&(p<=ap[:,2]),(bp[:,1]<=p)&(p<=bp[:,2])),
                              'width90':(ap[:,2]-ap[:,1],bp[:,2]-bp[:,1])}
                    for name,(x,y) in measures.items():
                        tables['paired_efficacy'].append({**tk,'basket':j+1,'comparison':'Proposed minus '+method,
                                   'metric':name,**paired_fields(x,y),'reps':R})
        if not config['cluster_only']:
            if e.family=='main' and e.dynamics=='homogeneous':
                baseline_time[(e.scenario,e.n)]=(post.copy(),true_p.copy())
            elif e.family=='sensitivity' and (e.scenario,e.n) in baseline_time:
                bp,bptrue=baseline_time[(e.scenario,e.n)]
                for m,method in enumerate(methods):
                    for j in range(5):
                        x=(post[:,m,j,1]>e.p0).astype(float);y=(bp[:,m,j,1]>e.p0).astype(float)
                        avg,se=mean_se(x);bavg,bse=mean_se(y);diff=avg-bavg;dse=math.sqrt(se**2+bse**2)
                        tables['dynamics_comparison'].append({**tk,'method':method,'basket':j+1,
                           'rejection_difference':diff,'mcse':dse,'lo95':diff-Z95*dse,'hi95':diff+Z95*dse,
                           'true_ORR_change':float(true_p[j]-bptrue[j]),
                           'comparison':'Time-varying minus homogeneous; independent per-setting streams'})
        print('Summarized '+e.name,flush=True)
    for name,rows in tables.items():write_csv(out/(name+'.csv'),rows)
    if excel:export_excel(out/'results.xlsx',tables,manifest)
    write_tex(out/'cluster_tables.tex',tables['cluster_metrics'],exps,methods,'clusters')
    if tables['basket_metrics']:write_tex(out/'efficacy_tables.tex',tables['basket_metrics'],exps,methods,'efficacy')
    note=('Full ORR inference run.' if not config['cluster_only'] else 'CLUSTER-ONLY CHECK: no ORR inferential results.')
    (out/'INTERPRETATION.txt').write_text(note+'\nProposed = frozen Hitting_mix50, not original silhouette.\n'
        'The partition is selected from observed trajectories only. Stage 2 uses actual binary counts.\n'
        'Hierarchical hyperparameters are integrated numerically; partition uncertainty is NOT integrated.\n'
        'A 90% credible interval rule is not calibrated 5% type I error control.\n'
        'Cluster count accuracy and exact membership accuracy are different metrics.\n'
        'Average credible interval endpoints are NOT a confidence interval for a Monte Carlo mean.\n'
        'Across methods, differences are paired within the same simulated trial. Across dynamics, streams are independent.\n'
        'The stress homogeneity settings can be efficacy alternatives; inspect truth.csv.\n',encoding='utf-8')
    return tables


def export_excel(path,tables,manifest):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook();ws=wb.active;ws.title='Read_me'
    notes=[('Item','Value'),('Suite version',manifest['version']),('Primary','Proposed = frozen Hitting_mix50'),
           ('Purpose','CLUSTER-ONLY DIAGNOSTIC' if manifest['config']['cluster_only'] else 'Full cluster selection and ORR posterior evaluation'),
           ('Replicates per setting',manifest['config']['reps']),('Seed',manifest['config']['seed']),
           ('Posterior','Numerical integration over mu/tau, conditional on the selected partition'),
           ('Error control','Not calibrated. Report empirical type I error and power, not nominal guarantees.'),
           ('Units','Probabilities shown as percentages; raw CSV probabilities are between zero and one. MCSE refers to simulation error.'),
           ('Source choice',manifest['config']['supp_source']),('Threshold mode',manifest['config']['threshold_mode']),
           ('Data source','Generated from settings recorded in manifest.json; frozen module hashes in provenance/frozen_sources.json'),
           ('No oracle inputs','Observed trajectories only are passed to analysis. True parameters appear only in generation and evaluation.'),
           ('Cluster accuracy','Correct number is not the same as correct partition. See both columns.'),
           ('Method details','METHODS.md; ORR-only is NOT modified by the new refinement.'),
           ('SciPy documentation','https://docs.scipy.org/doc/scipy/reference/generated/scipy.integrate.cumulative_simpson.html')]
    for r in notes:ws.append(r)
    for name,rows in tables.items():
        if not rows:continue
        sh=wb.create_sheet(name[:31]);fields=list(dict.fromkeys(k for r in rows for k in r));sh.append(fields)
        for row in rows:
            sh.append([None if isinstance(row.get(k),float) and not math.isfinite(row[k]) else row.get(k) for k in fields])
    for sh in wb:
        sh.sheet_view.showGridLines=False;sh.freeze_panes='F2' if sh.title!='Read_me' else 'A2';sh.auto_filter.ref=sh.dimensions
        sh.row_dimensions[1].height=48
        for c in sh[1]:
            c.font=Font(name='Calibri',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
            c.alignment=Alignment(wrap_text=True,vertical='center')
        for col in sh.columns:
            name=str(col[0].value)
            width={'setting':51,'description':96,'source':25,'method':25,'comparison':37,'error_control':52,
                   'decision_interpretation':20,'efficacy_status':20,'dynamics':23,'family':16}.get(name,min(32,max(13,len(name)+1)))
            sh.column_dimensions[col[0].column_letter].width=width
            for c in col[1:]:
                c.font=Font(name='Calibri',size=11,color='008000' if isinstance(c.value,(int,float)) else '666666')
                if isinstance(c.value,float):
                    rate_prefixes=('correct_number_','correct_partition_','under_','over_',
                                   'rejection_','coverage90_','FWER_','all_active_detected_')
                    percentage=(name.startswith('prob_K') or name in
                                ('probability','true_ORR','null_threshold','legacy_threshold','p0') or
                                (name.startswith(rate_prefixes) and name.endswith(('_rate','_lo95','_hi95','_mcse'))) or
                                name=='rejection_at_legacy_threshold')
                    c.number_format='0.00%' if percentage else '0.000000'
                if isinstance(c.value,str) and len(c.value)>width:
                    c.alignment=Alignment(wrap_text=True,vertical='top')
                    sh.row_dimensions[c.row].height=max(sh.row_dimensions[c.row].height or 15,15*(1+len(c.value)//int(width)))
        if sh.title=='Read_me':
            sh.column_dimensions['A'].width=26;sh.column_dimensions['B'].width=110
            for row in sh.iter_rows(min_row=2):
                row[1].alignment=Alignment(wrap_text=True,vertical='top');sh.row_dimensions[row[0].row].height=31
    wb.save(path)


def tex_escape(s):
    return str(s).replace('\\',r'\textbackslash{}').replace('_',r'\_').replace('%',r'\%').replace('&',r'\&')


def write_tex(path,rows,exps,methods,kind):
    lines=['% Generated from completed run; requires booktabs. Rates are percentages.',
           '% Proposed = Hitting_mix50. No claim of calibrated efficacy type I error.']
    for e in exps:
        lines += [r'\begin{table}[ht]',r'\centering',r'\caption{'+tex_escape(e.name)+'.}']
        if kind=='clusters':
            lines += [r'\begin{tabular}{lrrr}',r'\toprule',r'Method & Correct number (\%) & Exact partition (\%) & Mean ARI \\',r'\midrule']
            for m in methods:
                rr=[r for r in rows if r['setting']==e.name and r['method']==m]
                if rr:
                    r=rr[0];lines.append(tex_escape(m)+f" & {100*r['correct_number_rate']:.2f} & {100*r['correct_partition_rate']:.2f} & {r['ARI_mean']:.3f}"+r' \\')
        else:
            lines += [r'\begin{tabular}{lrrrrr}',r'\toprule',r'Method & Basket 1 & Basket 2 & Basket 3 & Basket 4 & Basket 5 \\',r'\midrule']
            for m in methods:
                rr=sorted([r for r in rows if r['setting']==e.name and r['method']==m],key=lambda r:r['basket'])
                if rr:lines.append(tex_escape(m)+' & '+' & '.join(f"{100*r['rejection_rate']:.2f}" for r in rr)+r' \\')
        lines += [r'\bottomrule',r'\end{tabular}',r'\end{table}']
    Path(path).write_text('\n'.join(lines)+'\n',encoding='utf-8')
