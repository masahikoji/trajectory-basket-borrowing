"""Summarize precision-prior sensitivity on matched saved trials."""
from __future__ import annotations
import csv,json,math
from pathlib import Path
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from posterior import independent_table


def mean_se(x):
    x=np.asarray(x,float)
    return float(x.mean()),float(x.std(ddof=1)/math.sqrt(len(x)))


def write_csv(path,rows):
    if not rows:return
    with Path(path).open('w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)


def assemble(a,e,methods,lookup):
    ids=a['request_ids'];pos=a['positions']
    out=np.empty(ids.shape+(7,))
    for m,method in enumerate(methods):
        if method=='Independent':
            out[:,m]=independent_table(e['n'],e['p0'],e['legacy_p0'])[a['responses']]
        else:
            if np.any(ids[:,m]<0):raise AssertionError('Missing posterior request.')
            out[:,m]=lookup[ids[:,m],pos[:,m]]
    if not np.isfinite(out[...,:6]).all():raise ArithmeticError('Nonfinite posterior output.')
    return out


def replay_check(actual,reference,p0,methods,tolerance=2e-5):
    rows=[]
    for m,method in enumerate(methods):
        columns=7 if method=='EXNEX' else 6
        ref=reference[:,m,:,:columns];now=actual[:,m,:,:columns]
        available=np.isfinite(ref).all()
        delta=float(np.max(abs(now-ref))) if available else None
        flips=int(np.count_nonzero((actual[:,m,:,1]>p0)!=(reference[:,m,:,1]>p0))) if available else None
        rows.append(dict(method=method,available=bool(available),max_abs_difference=delta,
                         rejection_disagreements=flips,passed=(delta<=tolerance) if available else None))
    return rows


def summarize(out,metadata,lookups,excel=True):
    out=Path(out);methods=metadata['methods'];rates=metadata['rates'];R=metadata['used_reps']
    tables={s:[] for s in ('prior_specs','basket_metrics','paired_priors','trial_metrics','primary_replay','input_audit')}
    for b in rates:
        tables['prior_specs'].append(dict(shape=2,rate=b,E_precision=2/b,E_logodds_variance=b,
                    primary=(b==1),mu_mean=0,mu_variance=1,
                    interpretation='Smaller prior heterogeneity scale' if b<1 else ('Primary' if b==1 else 'Larger prior heterogeneity scale')))
    for e in metadata['settings']:
        with np.load(out/'inputs'/f'{e["name"]}.npz',allow_pickle=False) as z: a={k:z[k] for k in z.files}
        p=np.asarray(e['true_orr']);null=p<=e['p0']+1e-10
        common=dict(setting=e['name'],family=e['family'],scenario=e['scenario'],n=e['n'],reps=R)
        outputs={b:assemble(a,e,methods,lookups[b]) for b in rates}
        base=outputs[1.0]
        for row in replay_check(base,a['source_post'],e['p0'],methods):
            tables['primary_replay'].append({**common,**row})
        tables['input_audit'].append({**common,'source':e['source_root'],
                    'response_sha256':e['response_hash'],'partition_sha256':e['partition_hash'],
                    'counts_and_partitions_changed':False,'p0':e['p0'],'legacy_p0':e['legacy_p0']})
        for b,post in outputs.items():
            for m,method in enumerate(methods):
                v=post[:,m];ref=base[:,m];reject=v[:,:,1]>e['p0'];refrej=ref[:,:,1]>e['p0']
                for j in range(5):
                    mean,mse=mean_se(v[:,j,0]);width,wse=mean_se(v[:,j,2]-v[:,j,1]);rej,rse=mean_se(reject[:,j])
                    coverage,cse=mean_se((v[:,j,1]<=p[j])&(p[j]<=v[:,j,2]))
                    row={**common,'method':method,'basket':j+1,'shape':2,'rate':b,'true_orr':float(p[j]),'p0':e['p0'],
                         'decision_interpretation':'type_I_error' if null[j] else 'power',
                         'rejection_rate':rej,'rejection_mcse':rse,
                         'posterior_mean':mean,'posterior_mean_mcse':mse,
                         'mean_lower90':float(v[:,j,1].mean()),'mean_upper90':float(v[:,j,2].mean()),
                         'mean_width90':width,'width_mcse':wse,'coverage90':coverage,'coverage_mcse':cse,
                         'cdf_ci_disagreements':int(np.count_nonzero(reject[:,j]!=(v[:,j,3]>.95)))}
                    tables['basket_metrics'].append(row)
                    if b!=1:
                        for label,x,y in [('rejection',reject[:,j],refrej[:,j]),('posterior_mean',v[:,j,0],ref[:,j,0]),
                                           ('width90',v[:,j,2]-v[:,j,1],ref[:,j,2]-ref[:,j,1])]:
                            delta,se=mean_se(np.asarray(x,float)-np.asarray(y,float))
                            tables['paired_priors'].append({**common,'method':method,'basket':j+1,'rate':b,'reference_rate':1,
                                  'metric':label,'difference':delta,'mcse':se,'lo95':delta-1.95996398454*se,'hi95':delta+1.95996398454*se})
                fwer,fse=mean_se(np.any(reject[:,null],axis=1)) if null.any() else (None,None)
                tables['trial_metrics'].append({**common,'method':method,'rate':b,'null_baskets':int(null.sum()),
                                                'fwer':fwer,'fwer_mcse':fse,'calibration':'No new cutoff calibration'})
            target=out/'posteriors'/f'{e["name"]}_rate{b:g}.npz';target.parent.mkdir(exist_ok=True)
            from prior_io import atomic_npz
            atomic_npz(target,post=post,replicate=a['replicate'],responses=a['responses'],labels=a['labels'],methods=np.array(methods),rate=b)
        print(f'Summarized {e["name"]}',flush=True)
    for name,rows in tables.items():write_csv(out/f'{name}.csv',rows)
    failed=[r for r in tables['primary_replay'] if r['available'] and not r['passed']]
    check=dict(primary_replay_passed=not failed,baseline_comparisons=len(tables['primary_replay']),
               unavailable=sum(not r['available'] for r in tables['primary_replay']),
               max_abs_difference=max((r['max_abs_difference'] or 0) for r in tables['primary_replay']),
               rejection_disagreements=sum(r['rejection_disagreements'] or 0 for r in tables['primary_replay']),
               note='Only numerical reproducibility of the primary prior, not clinical/statistical robustness.')
    from posterior import atomic_json
    atomic_json(out/'primary_replay_summary.json',check)
    if failed:raise ArithmeticError('Primary-prior replay differs from stored summaries by more than 2e-5. See primary_replay.csv. Do not interpret this as prior sensitivity until the difference is resolved.')
    if excel:make_excel(out/'results_prior.xlsx',tables,metadata)
    write_tex(out/'prior_sensitivity_tables.tex',tables['basket_metrics'],metadata)
    return check


def make_excel(path,tables,metadata):
    wb=Workbook();wb.remove(wb.active)
    intro=wb.create_sheet('Read_me');intro.sheet_view.showGridLines=False
    notes=[('Prior sensitivity', 'SAVED DATA REPLAY - no new trajectories or clustering'),
           ('Primary method',metadata.get('primary','Hitting_mix50')),
           ('Trials per setting',metadata['used_reps']),('Settings',len(metadata['settings'])),
           ('Priors','Gamma(shape=2, rate=0.5 / 1 / 2); mu ~ N(0,1) fixed'),
           ('Decision','Equal-tail 90% posterior interval lower endpoint > source p0. No 5% calibration.'),
           ('Scope','Main S1-S3 at n=20/30/50; supplementary S1-S3 at n=30. No crossing with time/stress/mixture-weight settings by default.'),
           ('Sources','Unchanged source manifest and raw trial counts/partitions; see input_audit and manifest.json'),
           ('Uncertainty','MCSE describes simulation variation. Prior comparisons are paired. These are not simultaneous confidence statements.'),
           ('Independent','Beta(1,1) is unchanged. EXNEX: only the EX precision prior changes; EX weight and NEX prior are held fixed.'),
           ('Caution','SMOKE CHECK ONLY - NOT FOR MANUSCRIPT' if metadata['smoke'] else 'Review rejection, interval-width and replay checks before writing conclusions.'),
           ('Partition uncertainty','Second-stage inference is conditional on saved partitions; partition uncertainty is not marginalized.')]
    for row in notes:intro.append(row)
    intro.column_dimensions['A'].width=24;intro.column_dimensions['B'].width=105
    for row in intro:
        intro.row_dimensions[row[0].row].height=38
        for c in row:c.alignment=Alignment(vertical='top',wrap_text=True);c.font=Font(name='Calibri',size=11,color='555555')
        row[0].font=Font(name='Calibri',size=11,bold=True,color='16324F')
    intro['B11'].font=Font(name='Calibri',size=11,bold=True,color='BF6900')
    all_tables={'Priors':tables['prior_specs'],
                'Main':[r for r in tables['basket_metrics'] if r['family']=='main'],
                'Supplement':[r for r in tables['basket_metrics'] if r['family']=='supplement'],
                'Paired':tables['paired_priors'],'Trial':tables['trial_metrics'],
                'Primary_replay':tables['primary_replay'],'Input_audit':tables['input_audit']}
    for name,rows in all_tables.items():
        ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False
        if not rows:continue
        keys=list(rows[0]);ws.append(keys);ws.freeze_panes='A2'
        for r in rows:ws.append([r[k] for k in keys])
        ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=32
        for c in ws[1]:c.fill=PatternFill('solid',fgColor='16324F');c.font=Font(name='Calibri',bold=True,color='FFFFFF',size=10);c.alignment=Alignment(wrap_text=True,vertical='center')
        for col,key in enumerate(keys,1):
            width=43 if key=='setting' else 22 if key=='method' else 18
            if 'sha256' in key or key in ('source','interpretation'):width=36
            ws.column_dimensions[get_column_letter(col)].width=width
            for i in range(2,len(rows)+2):
                c=ws.cell(i,col);c.font=Font(name='Calibri',size=10,color='00804A' if isinstance(c.value,(float,int)) else '666666')
                c.alignment=Alignment(vertical='top',wrap_text=isinstance(c.value,str))
                if isinstance(c.value,float):c.number_format='0.0000'
                if key in ('rejection_rate','rejection_mcse','coverage90','coverage_mcse','fwer','fwer_mcse'):
                    c.number_format='0.0%'
                if i%2==0:c.fill=PatternFill('solid',fgColor='F2F6FA')
        if name=='Priors':
            for i in range(2,len(rows)+2):
                ws.cell(i,3,f'=A{i}/B{i}');ws.cell(i,4,f'=B{i}/(A{i}-1)')
                for j in (3,4):ws.cell(i,j).font=Font(name='Calibri',color='000000',size=10)
        if name in ('Main','Supplement'):
            l=keys.index('mean_lower90')+1;u=keys.index('mean_upper90')+1;j=keys.index('mean_width90')+1
            for i in range(2,len(rows)+2):
                ws.cell(i,j,f'={get_column_letter(u)}{i}-{get_column_letter(l)}{i}').font=Font(name='Calibri',size=10,color='000000')
    wb.save(path)


def write_tex(path,rows,meta):
    def escape(s):return str(s).replace('_',r'\_').replace('%',r'\%')
    lines=[r'% Requires \usepackage{booktabs}; generated from saved-trial posterior replay.',
           r'% Each row is basket-specific. Type I / power classification uses true ORR and the unchanged source threshold.',
           r'% Credible interval width is averaged over trials; not a confidence interval for a Monte Carlo mean.']
    if meta['smoke']:lines.append('% SMOKE CHECK ONLY - NOT FOR MANUSCRIPT')
    for family in ('main','supplement'):
        for n in sorted({r['n'] for r in rows if r['family']==family}):
            for scenario in ('1','2','3'):
                selected=[r for r in rows if r['family']==family and r['n']==n and str(r['scenario'])==scenario]
                if not selected:continue
                lines += [r'\begin{table}[htbp]',r'\centering\small',
                          rf'\caption{{Prior-scale sensitivity: {escape(family)}, scenario {scenario}, $n_j={n}$.}}',
                          r'\begin{tabular}{llrrrr}',r'\toprule',
                          r'Method & Basket/status & Rate & Reject (\%) & Post. mean & Width \\',r'\midrule']
                for r in selected:
                    status='I' if r['decision_interpretation']=='type_I_error' else 'P'
                    lines.append(f'{escape(r["method"])} & {r["basket"]}/{status} & {r["rate"]:g} & {100*r["rejection_rate"]:.1f} & {r["posterior_mean"]:.3f} & {r["mean_width90"]:.3f} '+r'\\')
                lines += [r'\bottomrule\end{tabular}',r'\par\footnotesize I: type I error; P: power. Gamma shape--rate priors have shape 2. Partitions and response counts are identical across prior settings.',r'\end{table}']
    Path(path).write_text('\n'.join(lines)+'\n')
