"""Transition-estimation summaries and publication-oriented tables."""
from __future__ import annotations
import csv
import json
from pathlib import Path
import math
import os
import zipfile
from xml.etree import ElementTree as ET
import numpy as np
from transition_model import (STATES, OBS_PROB, FIELDS, theoretical_support,
                              estimate_transition, validate_arrays)


SUPPORT_BANDS = ((0, 0, '0'), (1, 4, '1-4'), (5, 9, '5-9'),
                 (10, 19, '10-19'), (20, None, '20+'))


def mean_se(x):
    x = np.asarray(x, dtype=float)
    if x.size == 0:
        return float('nan'), float('nan')
    return float(x.mean()), float(x.std(ddof=1)/np.sqrt(x.size)) if x.size > 1 else float('nan')


def estimate_error_stats(estimates, truth):
    estimates = np.asarray(estimates, dtype=float)
    N = len(estimates)
    keys = ('mean_estimate', 'bias', 'empirical_sd', 'mse', 'rmse', 'mae',
            'mcse_bias', 'mcse_mse', 'mcse_rmse_delta', 'mcse_mae',
            'estimate_q05', 'estimate_median', 'estimate_q95',
            'abs_error_q95', 'prob_abs_error_gt_0p10', 'prob_abs_error_gt_0p20')
    if N == 0:
        return {k: float('nan') for k in keys}
    error = estimates - truth
    mse, mse_se = mean_se(error**2)
    bias, bias_se = mean_se(error)
    mae, mae_se = mean_se(np.abs(error))
    rmse = np.sqrt(mse)
    qs = np.quantile(estimates, [.05, .5, .95])
    return dict(mean_estimate=float(estimates.mean()), bias=bias,
                empirical_sd=float(estimates.std(ddof=1)) if N > 1 else float('nan'),
                mse=mse, rmse=float(rmse), mae=mae, mcse_bias=bias_se,
                mcse_mse=mse_se, mcse_rmse_delta=float(mse_se/(2*rmse)) if rmse > 0 else float('nan'),
                mcse_mae=mae_se, estimate_q05=float(qs[0]), estimate_median=float(qs[1]),
                estimate_q95=float(qs[2]), abs_error_q95=float(np.quantile(np.abs(error), .95)),
                prob_abs_error_gt_0p10=float(np.mean(np.abs(error) > .10)),
                prob_abs_error_gt_0p20=float(np.mean(np.abs(error) > .20)))


def row_error_stats(errors):
    errors = np.asarray(errors, dtype=float)
    keys = ('row_mse', 'row_rmse', 'mean_total_variation', 'tv_q95',
            'mean_max_abs_error', 'max_abs_error_q95', 'mcse_row_mse',
            'mcse_row_rmse_delta', 'mcse_mean_tv')
    if len(errors) == 0:
        return {k: float('nan') for k in keys}
    squared = np.mean(errors**2, axis=-1)
    tv = .5*np.abs(errors).sum(-1)
    maxabs = np.max(np.abs(errors), axis=-1)
    mse, se_mse = mean_se(squared)
    mean_tv, se_tv = mean_se(tv)
    rmse = math.sqrt(mse)
    return dict(row_mse=mse, row_rmse=rmse, mean_total_variation=mean_tv,
                tv_q95=float(np.quantile(tv, .95)), mean_max_abs_error=float(maxabs.mean()),
                max_abs_error_q95=float(np.quantile(maxabs, .95)), mcse_row_mse=se_mse,
                mcse_row_rmse_delta=se_mse/(2*rmse) if rmse > 0 else float('nan'),
                mcse_mean_tv=se_tv)


def write_csv(path, rows):
    if not rows:
        return
    with Path(path).open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        for row in rows:
            writer.writerow({k: '' if isinstance(v, (float, np.floating)) and not np.isfinite(v) else v
                             for k, v in row.items()})


def load_arrays(out, e, reps, chunk_size, fingerprint):
    values = {k: [] for k in FIELDS}
    for start in range(0, reps, chunk_size):
        stop = min(start+chunk_size, reps)
        file = out/'raw'/f'{e.name}_r{start:07d}_{stop:07d}.npz'
        with np.load(file, allow_pickle=False) as z:
            if str(z['fingerprint']) != fingerprint or str(z['experiment']) != e.name:
                raise ValueError(f'Invalid checkpoint: {file}')
            if not np.array_equal(z['replicate'], np.arange(start, stop)):
                raise ValueError(f'Incomplete trial sequence: {file}')
            part = {k: z[k] for k in FIELDS}
            validate_arrays(part, e.n, stop-start)
            for k in FIELDS:
                values[k].append(part[k])
    return {k: np.concatenate(v) for k, v in values.items()}


def summarize_setting(e, arrays):
    R = len(arrays['Phat']); fitted = estimate_transition(arrays['transition_counts'])
    truth = e.truth
    outgoing = arrays['transition_counts'].sum(-1)
    visits = outgoing + arrays['final_counts']
    base = dict(setting=e.name, scenario=e.scenario, n=e.n)
    cells = []; rows = []; strata = []; horizons = []; matrices = []
    decomposition = []
    for j in range(5):
        key = dict(base, basket=j+1, mechanism=chr(65+int(e.mechanisms[j])))
        theory = theoretical_support(e.initial[j], truth[j], e.n)
        total_visits = arrays['length_counts'][:, j] @ np.arange(1, 11)
        mean_horizon = total_visits/e.n
        horizon_mean, horizon_se = mean_se(mean_horizon)
        horizons.append(dict(key, trials=R, mean_assessments_per_patient=horizon_mean,
                        mcse_mean_assessments=horizon_se,
                        q05_trial_mean_assessments=float(np.quantile(mean_horizon,.05)),
                        q95_trial_mean_assessments=float(np.quantile(mean_horizon,.95)),
                        expected_assessments_per_patient=float(OBS_PROB @ np.arange(1,11)),
                        mean_total_transitions=float(outgoing[:,j].sum(-1).mean())))
        matrix_err = fitted[:,j] - truth[j]
        matrix_mse, matrix_se = mean_se((matrix_err**2).mean(axis=(-1,-2)))
        matrices.append(dict(key, trials=R, matrix_mse=matrix_mse,
                        matrix_rmse=math.sqrt(matrix_mse), mcse_matrix_mse=matrix_se,
                        mean_frobenius_error=float(np.sqrt((matrix_err**2).sum((-1,-2))).mean()),
                        prob_any_empty_row=float(np.mean((outgoing[:,j]==0).any(-1)))))
        for r, origin in enumerate(STATES):
            rowkey = dict(key, from_state=origin)
            N = outgoing[:,j,r]; observed = N>0; empty = ~observed
            support_mean, support_se = mean_se(N)
            pempty, se_empty = mean_se(empty)
            cond_errors = matrix_err[observed,r]
            allstats = row_error_stats(matrix_err[:,r])
            condstats = row_error_stats(cond_errors)
            visit_prop = visits[:,j,r]/total_visits
            row = dict(rowkey, trials=R, nonempty_trials=int(observed.sum()),
                       mean_outgoing=support_mean, mcse_mean_outgoing=support_se,
                       expected_outgoing=float(theory['expected_outgoing'][r]),
                       outgoing_q05=float(np.quantile(N,.05)), outgoing_median=float(np.median(N)),
                       outgoing_q95=float(np.quantile(N,.95)),
                       empty_probability=pempty, mcse_empty_probability=se_empty,
                       exact_empty_probability=float(theory['empty_probability'][r]),
                       probability_outgoing_lt_5=float(np.mean(N<5)),
                       probability_outgoing_lt_10=float(np.mean(N<10)),
                       probability_outgoing_lt_20=float(np.mean(N<20)),
                       mean_observed_visits_in_state=float(visits[:,j,r].mean()),
                       mean_visit_weighted_state_fraction=float(visit_prop.mean()), **allstats)
            row.update({'nonempty_'+k: v for k,v in condstats.items()})
            for s, destination in enumerate(STATES):
                row['rmse_to_'+destination] = float(np.sqrt(np.mean(matrix_err[:,r,s]**2)))
                row['nonempty_rmse_to_'+destination] = (float(np.sqrt(np.mean(cond_errors[:,s]**2)))
                                                       if len(cond_errors) else float('nan'))
                empty_contribution = pempty*(.25-truth[j,r,s])**2
                nonempty_contribution = float(np.sum(matrix_err[observed,r,s]**2)/R)
                overall_mse = float(np.mean(matrix_err[:,r,s]**2))
                decomposition.append(dict(rowkey, to_state=destination, true_probability=float(truth[j,r,s]),
                            trials=R, empty_probability=pempty, mse_all=overall_mse,
                            empty_row_mse_contribution=float(empty_contribution),
                            nonempty_row_mse_contribution=nonempty_contribution,
                            decomposition_residual=float(overall_mse-empty_contribution-nonempty_contribution)))
                for scope, mask in (('all',np.ones(R,bool)), ('nonempty',observed)):
                    cells.append(dict(rowkey, to_state=destination, scope=scope,
                                      trials_total=R, trials_used=int(mask.sum()),
                                      true_probability=float(truth[j,r,s]),
                                      **estimate_error_stats(fitted[mask,j,r,s], truth[j,r,s])))
            rows.append(row)
            for lo, hi, name in SUPPORT_BANDS:
                mask=(N>=lo) if hi is None else ((N>=lo)&(N<=hi))
                strata.append(dict(rowkey, support_band=name, trials_total=R,
                                   trials_used=int(mask.sum()), fraction=float(mask.mean()),
                                   mean_outgoing=float(N[mask].mean()) if mask.any() else float('nan'),
                                   **row_error_stats(matrix_err[mask,r])))
    return dict(cell_errors=cells, row_errors=rows, support_strata=strata,
                observation_counts=horizons, matrix_errors=matrices, mse_decomposition=decomposition)


def cache_formulas(path, cached):
    ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    uri = ns['s']; temp = Path(path).with_suffix('.cached.xlsx')
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename in cached:
                root = ET.fromstring(data)
                for cell in root.findall('.//s:c', ns):
                    address = cell.attrib.get('r')
                    if address not in cached[info.filename]:
                        continue
                    value = cached[info.filename][address]
                    cell.attrib.pop('t', None)
                    v = cell.find('s:v', ns)
                    if v is None:
                        v = ET.SubElement(cell, '{'+uri+'}v')
                    if value is None or not math.isfinite(float(value)):
                        cell.set('t','str'); v.text=''
                    else:
                        v.text=repr(float(value))
                data=ET.tostring(root, encoding='utf-8', xml_declaration=True)
            dst.writestr(info, data)
    os.replace(temp, path)


def write_workbook(path, tables, smoke=False):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.properties import CalcProperties
    wb=Workbook(); info=wb.active; info.title='Read_me'; info.sheet_view.showGridLines=False
    info.column_dimensions['A'].width=30; info.column_dimensions['B'].width=112
    details=[('TRANSITION PRECISION', 'Main simulations: S1-S3; n=20,30,50.'),
             ('Status', 'SOFTWARE CHECK ONLY' if smoke else 'Monte Carlo transition-estimation diagnostics'),
             ('Estimator','Observed row ratios; zero-outgoing rows are assigned 0.25 in every column.'),
             ('Primary evaluation','All trials, including zero-outgoing rows.'),
             ('Conditional evaluation','Nonempty rows only. This is conditional precision, not a replacement for the primary evaluation.'),
             ('Row RMSE','sqrt(mean over trials and the four destinations of squared error). Equal row weights; not exposure-weighted.'),
             ('TV','Half the row L1 error. Range 0 to 1. Q95 is a simulation distribution quantile, not a confidence bound.'),
             ('MCSE','Monte Carlo standard error across independent simulated trials. RMSE MCSE uses the delta method.'),
             ('Error thresholds','.10 and .20 are descriptive absolute probability errors, not established clinical adequacy cutoffs.'),
             ('Scope','Primary assessment-count law is fixed. Short/long follow-up, alternative smoothing, and ORR posterior performance are not rerun.'),
             ('References','manifest.json, execution.json, METHODS.md, and raw/*.npz in the same results folder.'),
             ('Source','Original study mechanisms and observation model; all statistics computed from observed counts.')]
    for ri,values in enumerate(details,1):
        info.append(values); info.row_dimensions[ri].height=38
        for c in info[ri]: c.alignment=Alignment(wrap_text=True,vertical='center')
        info.cell(ri,1).font=Font(name='Calibri',bold=True,color='455A64')
    for c in info[1]:
        c.fill=PatternFill('solid',fgColor='17365D'); c.font=Font(name='Calibri',bold=True,color='FFFFFF',size=14)
    caches={}
    font_imported=Font(name='Calibri',size=10,color='00804A')
    font_formula=Font(name='Calibri',size=10,color='000000')
    font_static=Font(name='Calibri',size=10,color='666666')
    font_header=Font(name='Calibri',bold=True,color='FFFFFF',size=10)
    align_body=Alignment(vertical='center')
    align_header=Alignment(wrap_text=True,vertical='center')
    fill_header=PatternFill('solid',fgColor='17365D')
    def add_sheet(name, records, columns=None):
        if not records: return
        ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False
        columns=columns or list(records[0]); ws.append(columns)
        indices={k:i+1 for i,k in enumerate(columns)}; expected={}
        formula_specs={'bias':('mean_estimate','true_probability','sub'),
                       'rmse':('mse',None,'sqrt'), 'row_rmse':('row_mse',None,'sqrt'),
                       'nonempty_row_rmse':('nonempty_row_mse',None,'sqrt'),
                       'matrix_rmse':('matrix_mse',None,'sqrt')}
        for ri, record in enumerate(records,2):
            for ci,key in enumerate(columns,1):
                value=record.get(key)
                if isinstance(value,(float,np.floating)) and not np.isfinite(value): value=None
                cell=ws.cell(ri,ci,value); cell.font=font_imported
                cell.alignment=align_body
                if key in formula_specs:
                    a,b,op=formula_specs[key]
                    if a in indices and (b is None or b in indices):
                        x=f'{get_column_letter(indices[a])}{ri}'
                        expr=(f'{x}-{get_column_letter(indices[b])}{ri}' if op=='sub' else f'SQRT({x})')
                        cell.value=f'=IF(ISNUMBER({x}),{expr},"")'
                        expected[cell.coordinate]=value
                        cell.font=font_formula
                if key in ('scenario','n','basket','trials','trials_used','trials_total','nonempty_trials'):
                    cell.number_format='0'
                elif 'probability' in key or key.startswith('prob_') or key=='fraction':
                    cell.number_format='0.0%'
                elif isinstance(value,(int,float,np.integer,np.floating)):
                    cell.number_format='0.0000' if ('mse' in key or 'bias' in key or 'error' in key) else '0.000'
                else:
                    cell.font=font_static
            ws.row_dimensions[ri].height=17
        for ci,key in enumerate(columns,1):
            c=ws.cell(1,ci);c.value=key.replace('_','\n');c.font=font_header
            c.fill=fill_header;c.alignment=align_header
            ws.column_dimensions[get_column_letter(ci)].width=(42 if key=='setting' else 16 if len(key)>12 else 12)
        ws.row_dimensions[1].height=76;ws.freeze_panes='F2' if len(columns)>5 else 'A2'
        ws.auto_filter.ref=ws.dimensions
        index=wb.sheetnames.index(name)+1
        if expected:caches[f'xl/worksheets/sheet{index}.xml']=expected
    summary_cols=['scenario','n','basket','mechanism','from_state','mean_outgoing','empty_probability',
                  'row_rmse','nonempty_row_rmse','mean_total_variation','tv_q95',
                  'rmse_to_CR','rmse_to_PR','rmse_to_SD','rmse_to_PD']
    for n in sorted({r['n'] for r in tables['row_errors']}):
        add_sheet(f'Main_n{n}',[r for r in tables['row_errors'] if r['n']==n],summary_cols)
    add_sheet('Cell_errors',tables['cell_errors'])
    add_sheet('Row_errors',tables['row_errors'])
    add_sheet('Support_strata',tables['support_strata'])
    add_sheet('MSE_decomposition',tables['mse_decomposition'])
    add_sheet('Observation_counts',tables['observation_counts'])
    add_sheet('Matrix_errors',tables['matrix_errors'])
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=False)
    wb.save(path);cache_formulas(path,caches)
    check=load_workbook(path,data_only=True,read_only=False)
    for filename,entries in caches.items():
        idx=int(filename.split('sheet')[-1].split('.')[0])-1
        ws=check[check.sheetnames[idx]]
        for addr,value in entries.items():
            actual=ws[addr].value
            if value is not None and math.isfinite(float(value)) and not math.isclose(actual,float(value),rel_tol=1e-12,abs_tol=1e-15):
                raise ValueError(f'Excel cache mismatch: {ws.title}!{addr}')
    check.close()


def write_latex(path, rows):
    fmt=lambda x: '--' if not np.isfinite(x) else f'{x:.3f}'
    text=[r'\begin{longtable}{lllr rrrrr}',
          r'\caption{Transition-row estimation precision in the main simulations.}\label{tab:transition_precision}\\',
          r'\toprule Scenario & $n$ & Basket & Origin & Mean $N$ & Empty (\%) & RMSE & RMSE$_{N>0}$ & Mean TV \\\midrule',
          r'\endfirsthead',r'\toprule Scenario & $n$ & Basket & Origin & Mean $N$ & Empty (\%) & RMSE & RMSE$_{N>0}$ & Mean TV \\\midrule',
          r'\endhead',r'\bottomrule\endfoot']
    for r in rows:
        text.append(f"S{r['scenario']} & {r['n']} & {r['basket']} & {r['from_state']} & "
                    f"{r['mean_outgoing']:.1f} & {100*r['empty_probability']:.2f} & "
                    f"{fmt(r['row_rmse'])} & {fmt(r['nonempty_row_rmse'])} & {fmt(r['mean_total_variation'])}" + r' \\')
    text.extend([r'\end{longtable}',
                 r'\noindent\footnotesize RMSE averages squared probability error equally over the four destinations before taking a square root. '
                 r'The primary RMSE includes all trials and the uniform replacement for empty rows; RMSE$_{N>0}$ conditions on a nonempty outgoing row. '
                 r'TV is one-half the row $L_1$ error. These results describe estimation under the simulated mechanisms, not a general precision guarantee.'])
    Path(path).write_text('\n'.join(text)+'\n')


def export_reports(out, settings, reps, chunk_size, fingerprint, excel=True):
    out=Path(out)
    names=('cell_errors','row_errors','support_strata','observation_counts','matrix_errors','mse_decomposition')
    tables={k:[] for k in names}
    for e in settings:
        print(f'Summarizing {e.name}...', flush=True)
        arrays=load_arrays(out,e,reps,chunk_size,fingerprint)
        result=summarize_setting(e,arrays)
        for k in names: tables[k].extend(result[k])
    for k,rows in tables.items():write_csv(out/(k+'.csv'),rows)
    write_csv(out/'true_transitions.csv',[dict(setting=e.name,scenario=e.scenario,n=e.n,basket=j+1,
            mechanism=chr(65+int(e.mechanisms[j])),from_state=STATES[r],to_state=STATES[s],true_probability=float(e.truth[j,r,s]))
            for e in settings for j in range(5) for r in range(4) for s in range(4)])
    residual=max(abs(r['decomposition_residual']) for r in tables['mse_decomposition'])
    if residual>1e-12: raise ValueError('MSE decomposition check failed.')
    (out/'checks.json').write_text(json.dumps(dict(cell_rows=len(tables['cell_errors']),
        row_rows=len(tables['row_errors']),max_mse_decomposition_residual=residual),indent=2)+'\n')
    manifest=json.loads((out/'manifest.json').read_text())
    if excel:write_workbook(out/'transition_precision.xlsx',tables,manifest['config']['smoke'])
    write_latex(out/'transition_precision_table.tex',tables['row_errors'])
    return tables


def main():
    import argparse
    p=argparse.ArgumentParser(description='Create the Excel workbook from completed transition diagnostic CSVs.')
    p.add_argument('--results',type=Path,required=True)
    p.add_argument('--from-raw',action='store_true')
    p.add_argument('--no-excel',action='store_true')
    args=p.parse_args();out=args.results.expanduser().resolve()
    if args.from_raw:
        from transition_model import main_settings
        mf=json.loads((out/'manifest.json').read_text()); c=mf['config']
        export_reports(out,main_settings(c['sizes'],c['scenarios']),c['reps'],c['chunk_size'],mf['fingerprint'],not args.no_excel)
        print('Summary tables completed.',flush=True)
        return
    names=('cell_errors','row_errors','support_strata','observation_counts','matrix_errors','mse_decomposition')
    tables={}
    integers={'scenario','n','basket','trials','trials_used','trials_total','nonempty_trials'}
    strings={'setting','mechanism','from_state','to_state','scope','support_band'}
    for name in names:
        with (out/(name+'.csv')).open(newline='',encoding='utf-8') as f:
            rows=list(csv.DictReader(f))
        for row in rows:
            for key,value in row.items():
                if key not in strings:
                    row[key]=(int(value) if key in integers else float(value)) if value else float('nan')
        tables[name]=rows
    mf=json.loads((out/'manifest.json').read_text())
    write_workbook(out/'transition_precision.xlsx',tables,mf['config']['smoke'])
    print('Excel workbook completed.',flush=True)


if __name__=='__main__':
    main()
